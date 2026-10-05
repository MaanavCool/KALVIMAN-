# excel_sync.py
# ════════════════════════════════════════════════════════════════
# Kalviman — Formal Excel workbook builder.
#
# Rebuilds Kalviman_Students_Live.xlsx after every registration.
# Three sheets:
#   1. "All Students"      — every answer, colour-banded by section,
#                            PLUS an ML ANALYSIS band (best-fit etc.)
#   2. "Analysis Reports"  — the full written ML report per student
#   3. "Summary"           — live counts + ML cluster breakdown
#
# Professional font (Arial) throughout. No fragile formulas — all
# counts are computed in Python, so the file always opens clean.
# ════════════════════════════════════════════════════════════════

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import EXCEL_PATH

FONT = "Arial"

# ── Column map: (Excel header, db column, section colour) ─────────────
COLUMNS = [
    ("ID",                      "id",                    "1A3A5C"),
    ("Submitted At",            "submitted_at",          "1A3A5C"),

    # Section 1 — Basic Profile
    ("Name",                    "name",                  "1D4ED8"),
    ("Age",                     "age",                   "1D4ED8"),
    ("Email",                   "email",                 "1D4ED8"),
    ("Phone",                   "phone",                 "1D4ED8"),
    ("City",                    "city",                  "1D4ED8"),
    ("Country",                 "country",               "1D4ED8"),
    ("School / College",        "school_college",        "1D4ED8"),
    ("Degree Level",            "current_degree",        "1D4ED8"),
    ("Institution",             "institution_name",      "1D4ED8"),
    ("Academic Score",          "academic_percent",      "1D4ED8"),
    ("Parent Occupation",       "parent_occupation",     "1D4ED8"),

    # Section 2 — Self-Awareness
    ("Rate: Problem Solving",   "rate_problem_solving",  "D97706"),
    ("Rate: Creating",          "rate_creating",         "D97706"),
    ("Rate: Helping People",    "rate_helping_people",   "D97706"),
    ("Rate: Selling",           "rate_selling",          "D97706"),
    ("Rate: Data",              "rate_data",             "D97706"),
    ("Rate: Leading",           "rate_leading",          "D97706"),
    ("Rate: Building",          "rate_building",         "D97706"),
    ("Rate: Teaching",          "rate_teaching",         "D97706"),
    ("Interest Areas",          "interest_areas",        "D97706"),
    ("Flow Activity",           "activity_flow",         "D97706"),

    # Section 3 — Aptitude
    ("Problem Solving Style",   "problem_solving_style", "059669"),
    ("Learning Style",          "learning_style",        "059669"),
    ("Work Style",              "work_style",            "059669"),
    ("Decision Style",          "decision_style",        "059669"),
    ("Easy Subjects",           "easy_subjects",         "059669"),

    # Section 4 — Career
    ("Dream Careers",           "dream_careers",         "7C3AED"),
    ("Career Reason",           "career_reason",         "7C3AED"),
    ("Career Influence",        "career_influence",      "7C3AED"),
    ("Work Preference",         "work_preference",       "7C3AED"),
    ("Desired Impact",          "desired_impact",        "7C3AED"),

    # Section 5 — Lifestyle
    ("Income Priority",         "income_priority",       "DB2777"),
    ("Work Environment",        "work_environment",      "DB2777"),
    ("Preferred Country",       "preferred_country",     "DB2777"),

    # Section 6 — Education
    ("Education Budget",        "education_budget",      "0891B2"),
    ("Funding Comfort",         "funding_comfort",       "0891B2"),
    ("Languages",               "language_comfort",      "0891B2"),
    ("Willing to Relocate",     "willing_to_relocate",   "0891B2"),

    # Section 7 — Behaviour
    ("Rate: Finishes Tasks",    "rate_finishes_tasks",   "DC2626"),
    ("Rate: Works Pressure",    "rate_works_pressure",   "DC2626"),
    ("Rate: Asks Help",         "rate_asks_help",        "DC2626"),
    ("Rate: Recovers Setbacks", "rate_recovers_setbacks","DC2626"),
    ("Rate: Time Management",   "rate_time_management",  "DC2626"),
    ("Challenge Overcome",      "challenge_overcome",    "DC2626"),

    # Section 8 — Guidance
    ("Guidance Priority",       "guidance_priority",     "16A34A"),
    ("Clarity Needed",          "clarity_needed",        "16A34A"),
    ("Biggest Worry",           "biggest_worry",         "16A34A"),
    ("Success at Age 30",       "success_at_30",         "16A34A"),

    # ── ML ANALYSIS (new) ──
    ("ML Best-Fit Direction",   "ml_best_fit",           "0F766E"),
    ("Match %",                 "ml_match_percent",      "0F766E"),
    ("2nd Choice",              "_ml_second",            "0F766E"),
    ("3rd Choice",              "_ml_third",             "0F766E"),
    ("Resilience /5",           "ml_resilience",         "0F766E"),
]

SECTION_NAMES = {
    "1A3A5C": "SYSTEM",
    "1D4ED8": "SECTION 1 — BASIC PROFILE",
    "D97706": "SECTION 2 — SELF-AWARENESS",
    "059669": "SECTION 3 — APTITUDE",
    "7C3AED": "SECTION 4 — CAREER",
    "DB2777": "SECTION 5 — LIFESTYLE",
    "0891B2": "SECTION 6 — EDUCATION",
    "DC2626": "SECTION 7 — BEHAVIOUR",
    "16A34A": "SECTION 8 — GUIDANCE",
    "0F766E": "ML ANALYSIS (AUTO-GENERATED)",
}

WIDE_COLS = {
    "Name", "Email", "School / College", "Institution",
    "Dream Careers", "Career Reason", "Flow Activity", "Desired Impact",
    "Challenge Overcome", "Clarity Needed", "Biggest Worry", "Success at Age 30",
    "Easy Subjects", "ML Best-Fit Direction", "2nd Choice", "3rd Choice",
}

THIN = Side(style="thin", color="D5DBE2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _second_third(student, which):
    top3 = student.get("ml_top3_parsed") or []
    idx = 1 if which == "_ml_second" else 2
    if len(top3) > idx:
        name, pct = top3[idx]
        return f"{name} ({pct}%)"
    return ""


def rebuild_excel(students: list):
    wb = Workbook()

    # ══════════ Sheet 1: All Students ══════════
    ws = wb.active
    ws.title = "All Students"
    ws.freeze_panes = "C3"

    prev_color, sec_start = None, 1
    for ci, (header, _, color) in enumerate(COLUMNS, start=1):
        fill = PatternFill("solid", fgColor=color)
        h = ws.cell(row=2, column=ci, value=header)
        h.fill = fill
        h.font = Font(name=FONT, color="FFFFFF", bold=True, size=9)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = BORDER
        ws.cell(row=1, column=ci).fill = fill

        if color != prev_color:
            if prev_color is not None and ci - 1 >= sec_start:
                _band(ws, sec_start, ci - 1, prev_color)
            sec_start, prev_color = ci, color
    _band(ws, sec_start, len(COLUMNS), prev_color)

    # data rows
    for ri, s in enumerate(students, start=3):
        for ci, (header, db_col, _) in enumerate(COLUMNS, start=1):
            if db_col in ("_ml_second", "_ml_third"):
                val = _second_third(s, db_col)
            else:
                val = s.get(db_col, "")
            if val is None:
                val = ""
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = Font(name=FONT, size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            cell.border = BORDER
            if ri % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F7FA")

    for ci, (header, _, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = 36 if header in WIDE_COLS else 15
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 42
    # auto filter over the header row
    last_col = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A2:{last_col}2"

    # ══════════ Sheet 2: Analysis Reports ══════════
    ws2 = wb.create_sheet("Analysis Reports")
    headers = ["ID", "Name", "Best-Fit Direction", "Match %", "Resilience /5", "Full Analysis Report"]
    widths  = [10, 26, 34, 10, 13, 120]
    for ci, (htext, w) in enumerate(zip(headers, widths), start=1):
        c = ws2.cell(row=1, column=ci, value=htext)
        c.fill = PatternFill("solid", fgColor="0F2744")
        c.font = Font(name=FONT, color="FFFFFF", bold=True, size=10)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws2.column_dimensions[get_column_letter(ci)].width = w
    ws2.freeze_panes = "A2"
    ws2.row_dimensions[1].height = 24

    for ri, s in enumerate(students, start=2):
        gid = f"Kalviman-{str(s.get('id','')).zfill(5)}"
        ws2.cell(row=ri, column=1, value=gid).font = Font(name=FONT, bold=True, size=9, color="1D4ED8")
        ws2.cell(row=ri, column=2, value=s.get("name", "")).font = Font(name=FONT, bold=True, size=9)
        ws2.cell(row=ri, column=3, value=s.get("ml_best_fit", "")).font = Font(name=FONT, size=9)
        ws2.cell(row=ri, column=4, value=s.get("ml_match_percent", "")).font = Font(name=FONT, size=9)
        ws2.cell(row=ri, column=5, value=s.get("ml_resilience", "")).font = Font(name=FONT, size=9)
        rep = ws2.cell(row=ri, column=6, value=s.get("ml_report", "") or "")
        rep.font = Font(name=FONT, size=9)
        rep.alignment = Alignment(vertical="top", wrap_text=True)
        for ci in range(1, 6):
            ws2.cell(row=ri, column=ci).alignment = Alignment(vertical="top", wrap_text=True)
        # tall row so the wrapped report is readable
        ws2.row_dimensions[ri].height = 150

    # ══════════ Sheet 3: Summary ══════════
    ws3 = wb.create_sheet("Summary")
    ws3.column_dimensions["A"].width = 40
    ws3.column_dimensions["B"].width = 14
    title = ws3.cell(row=1, column=1, value="Kalviman — Live Summary")
    title.font = Font(name=FONT, bold=True, size=16, color="0F2744")

    from datetime import datetime
    ws3.cell(row=3, column=1, value="Total Students").font = Font(name=FONT, bold=True)
    ws3.cell(row=3, column=2, value=len(students))
    ws3.cell(row=4, column=1, value="Last Updated").font = Font(name=FONT, bold=True)
    ws3.cell(row=4, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    r = 6
    r = _count_block(ws3, r, "By ML Best-Fit Direction", students, "ml_best_fit")
    r = _count_block(ws3, r, "By Degree Level",          students, "current_degree")
    r = _count_block(ws3, r, "By Work Preference",       students, "work_preference")
    r = _count_block(ws3, r, "By Preferred Country",     students, "preferred_country")

    wb.save(EXCEL_PATH)
    print(f"Excel updated -> {EXCEL_PATH} ({len(students)} students)")


def _band(ws, c1, c2, color):
    ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
    cell = ws.cell(row=1, column=c1)
    cell.value = SECTION_NAMES.get(color, "")
    cell.font = Font(name=FONT, color="FFFFFF", bold=True, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _count_block(ws, start_row, title, students, col):
    counts = {}
    for s in students:
        k = s.get(col) or "Not specified"
        counts[k] = counts.get(k, 0) + 1
    t = ws.cell(row=start_row, column=1, value=title)
    t.font = Font(name=FONT, bold=True, size=11, color="0F2744")
    t.fill = PatternFill("solid", fgColor="EEF2F7")
    ws.cell(row=start_row, column=2).fill = PatternFill("solid", fgColor="EEF2F7")
    row = start_row + 1
    for k, v in sorted(counts.items(), key=lambda kv: (-kv[1], str(kv[0]))):
        ws.cell(row=row, column=1, value=k).font = Font(name=FONT, size=10)
        ws.cell(row=row, column=2, value=v).font = Font(name=FONT, size=10)
        row += 1
    return row + 2   # blank line before next block
